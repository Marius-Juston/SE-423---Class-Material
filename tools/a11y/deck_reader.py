"""Read-only access to the lecture decks for the accessibility tools.

Parsing goes through python-pptx (package, parts, relationships, slide/layout/master
inheritance, notes) and lxml XPath on its elements, never regexes over XML text.
Writing decks stays with tools/sync_deck_figures.py (Deck.save): python-pptx re-serializes
every part it saves, the faithful zip writer keeps everything else byte-identical.

Scripts importing this module need `python-pptx` and `coloraide` (and `pypandoc-binary`
for omml_to_mathml) in their inline script metadata.
"""

import functools
import io
import re
import tempfile
import zipfile
from pathlib import Path

from coloraide import Color
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "a16": "http://schemas.microsoft.com/office/drawing/2014/main",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main",
    "adec": "http://schemas.microsoft.com/office/drawing/2017/decorative",
}
UNKNOWN = "?"  # a colour we cannot reduce to one value (picture, gradient, pattern)
TITLE_TYPES = ("title", "ctrTitle")


def q(name):
    """'a:t' -> '{http://...drawingml...}t'."""
    prefix, local = name.split(":")
    return f"{{{NS[prefix]}}}{local}"


@functools.cache
def _compiled(path):
    return etree.XPath(path, namespaces=NS)


def xp(el, path):
    """XPath with the NS prefixes (python-pptx elements and plain lxml ones alike)."""
    return _compiled(path)(el)


def local(el):
    return etree.QName(el).localname


def open_deck(path):
    return Presentation(str(path))


def slides(prs):
    """(slide number, slide) for every slide, in order."""
    return list(enumerate(prs.slides, start=1))


def text_of(el):
    return "".join(t.text or "" for t in el.iter(q("a:t")))


def ph_type(sp):
    """Placeholder type of a p:sp ('body' when it is a placeholder without a type), or None."""
    ph = xp(sp, "./p:nvSpPr/p:nvPr/p:ph")
    return ph[0].get("type", "body") if ph else None


def title_shape(slide):
    """The slide's title placeholder element, or None."""
    for sp in xp(slide._element, ".//p:sp"):
        if ph_type(sp) in TITLE_TYPES:
            return sp
    return None


def title_text(slide):
    """Text of the slide's title placeholder ('' if empty), or None if there is none."""
    sp = title_shape(slide)
    return None if sp is None else text_of(sp).strip()


def cnvpr(el):
    """The p:cNvPr of a shape, picture, group or graphic frame."""
    found = xp(el, "./*[1]/p:cNvPr")
    return found[0] if found else None


def locator(el):
    """Stable id of a shape on its slide: its a16:creationId, else 'shape-<id>'."""
    c = cnvpr(el)
    cid = xp(c, ".//a16:creationId/@id") if c is not None else []
    return cid[0] if cid else f"shape-{c.get('id') if c is not None else '?'}"


def is_decorative(el):
    c = cnvpr(el)
    return c is not None and bool(xp(c, ".//adec:decorative[@val='1']"))


def is_media(pic):
    """Video/audio placeholder (its picture is only a poster frame)."""
    return bool(
        xp(
            pic,
            "./p:nvPicPr/p:nvPr/a:videoFile | ./p:nvPicPr/p:nvPr/a:audioFile | ./p:nvPicPr/p:nvPr/p:extLst//p14:media",
        )
    )


def rel_target(part, rid):
    """(target, is_external): an external URL, or the related part (None if missing)."""
    rel = part.rels.get(rid)
    if rel is None:
        return None, False
    if rel.is_external:
        return rel.target_ref, True
    return rel.target_part, False


def blob(part, rid):
    target, external = rel_target(part, rid)
    return None if external or target is None else target.blob


def picture_rid(pic, prefer_svg=False):
    """r:embed of the picture's image (its SVG when prefer_svg and there is one)."""
    if prefer_svg:
        svg = xp(pic, "./p:blipFill/a:blip//asvg:svgBlip/@r:embed")
        if svg:
            return svg[0]
    found = xp(pic, "./p:blipFill/a:blip/@r:embed")
    return found[0] if found else None


def notes_text(slide):
    """Speaker notes as a list of paragraphs."""
    if not slide.has_notes_slide:
        return []
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        return []
    return [p.text.strip() for p in tf.paragraphs if p.text.strip()]


def box(el):
    """(x, y, cx, cy) of a shape's own transform, or None."""
    xfrm = "./p:spPr/a:xfrm | ./p:grpSpPr/a:xfrm | ./p:xfrm"
    off = xp(el, f"({xfrm})/a:off")
    ext = xp(el, f"({xfrm})/a:ext")
    if not off or not ext:
        return None
    return int(off[0].get("x")), int(off[0].get("y")), int(ext[0].get("cx")), int(ext[0].get("cy"))


# ------------------------------------------------------------------ colours
def contrast(a, b):
    """WCAG 2.1 contrast ratio of two RRGGBB colours."""
    return Color(f"#{a}").contrast(f"#{b}", method="wcag21")


class Theme:
    """Resolves DrawingML colours (srgbClr / schemeClr / sysClr / prstClr with
    lumMod, lumOff, tint and shade) against a deck's theme and colour map."""

    PRESET = {"black": "000000", "white": "FFFFFF", "red": "FF0000"}

    def __init__(self, prs):
        master = prs.slide_master
        theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
        self.scheme = {}
        for el in xp(theme, ".//a:clrScheme/*"):
            val = xp(el, "./a:srgbClr/@val | ./a:sysClr/@lastClr")
            if val:
                self.scheme[local(el)] = val[0].upper()
        self.bg_styles = [local(el) for el in xp(theme, ".//a:bgFillStyleLst/*")]
        cmap = xp(master._element, "./p:clrMap")
        self.map = dict(cmap[0].attrib) if cmap else {}
        self.master_bg = self.background(master._element)

    def resolve(self, el):
        """el: one colour element (a:srgbClr, a:schemeClr, ...); returns RRGGBB or None."""
        kind, val = local(el), el.get("val")
        if kind == "srgbClr":
            rgb = val.upper()
        elif kind == "sysClr":
            rgb = (el.get("lastClr") or "").upper() or None
        elif kind == "prstClr":
            rgb = self.PRESET.get(val)
        elif kind == "schemeClr":
            rgb = self.scheme.get(self.map.get(val, val))
        else:
            rgb = None
        if not rgb:
            return None
        mods = {local(m): int(m.get("val")) / 100000 for m in el if local(m) in ("lumMod", "lumOff", "tint", "shade")}
        color = Color(f"#{rgb}")
        if "lumMod" in mods or "lumOff" in mods:
            hsl = color.convert("hsl")
            light = hsl["lightness"] * mods.get("lumMod", 1) + mods.get("lumOff", 0)
            color = hsl.set("lightness", min(max(light, 0), 1)).convert("srgb")
        r, g, b = color.convert("srgb").coords()
        if "tint" in mods:
            t = mods["tint"]
            r, g, b = (c * t + (1 - t) for c in (r, g, b))
        if "shade" in mods:
            t = mods["shade"]
            r, g, b = (c * t for c in (r, g, b))
        return "".join(f"{round(min(max(c, 0), 1) * 255):02X}" for c in (r, g, b))

    def solid(self, fill_parent):
        """Colour of the a:solidFill directly inside fill_parent: RRGGBB, UNKNOWN for a
        picture/gradient/pattern fill, None when there is no fill."""
        if xp(fill_parent, "./a:gradFill | ./a:blipFill | ./a:pattFill"):
            return UNKNOWN
        colour = xp(fill_parent, "./a:solidFill/*[1]")
        return (self.resolve(colour[0]) or UNKNOWN) if colour else None

    def background(self, root):
        """Background colour of a slide, layout or master (root element), or None."""
        bg = xp(root, "./p:cSld/p:bg")
        if not bg:
            return None
        ref = xp(bg[0], "./p:bgRef")
        if ref:
            i = int(ref[0].get("idx")) - 1001
            colour = xp(ref[0], "./*[1]")
            if 0 <= i < len(self.bg_styles) and self.bg_styles[i] == "solidFill" and colour:
                return self.resolve(colour[0])
            return UNKNOWN
        pr = xp(bg[0], "./p:bgPr")
        return self.solid(pr[0]) if pr else None

    def run_colour(self, rpr):
        """Text colour set on a run (a:rPr), or None when inherited."""
        if rpr is None:
            return None
        colour = xp(rpr, "./a:solidFill/*[1]")
        return self.resolve(colour[0]) if colour else None

    def layer(self, el):
        """(box, colour) for a shape or picture that hides what is behind it, else None."""
        b = box(el)
        if not b:
            return None
        if local(el) == "pic":
            return b, UNKNOWN
        sp_pr = xp(el, "./p:spPr")
        colour = self.solid(sp_pr[0]) if sp_pr else None
        return (b, colour) if colour else None

    def layers(self, root):
        """[(box, colour)] of filled shapes and pictures in z-order (back to front)."""
        return [lay for el in xp(root, ".//p:sp | .//p:pic") if (lay := self.layer(el))]


def background_under(b, layers, default):
    """Colour of the topmost filled shape containing the centre of box b."""
    cx, cy = b[0] + b[2] / 2, b[1] + b[3] / 2
    for (x, y, w, h), colour in reversed(layers):
        if x <= cx <= x + w and y <= cy <= y + h:
            return colour
    return default


# ------------------------------------------------------------------ equations
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CONTENT_TYPES = (
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>'
)
_ROOT_RELS = (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Target="word/document.xml" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"/></Relationships>'
)


def math_elements(el):
    """The m:oMathPara / m:oMath elements of the equations under el, in document order."""
    return xp(el, ".//a14:m/m:oMathPara | .//a14:m/m:oMath")


def omml_to_mathml(maths):
    """MathML for a list of OMML elements (m:oMath / m:oMathPara), with a TeX annotation.

    Uses pandoc's OMML reader (texmath, the converter behind pandoc's Word import): the
    equations are wrapped in a minimal .docx, one marked paragraph each, converted in a
    single pandoc call and split again. Returns one string per input ('' if pandoc
    produced nothing for it)."""
    if not maths:
        return []
    import pypandoc

    body = "".join(
        f"<w:p><w:r><w:t>@@EQ{i}@@</w:t></w:r>{etree.tostring(m).decode()}</w:p>" for i, m in enumerate(maths)
    )
    document = f'<w:document xmlns:w="{_W}" xmlns:m="{NS["m"]}"><w:body>{body}</w:body></w:document>'
    with tempfile.TemporaryDirectory() as tmp:
        docx = Path(tmp) / "equations.docx"
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", _CONTENT_TYPES)
            z.writestr("_rels/.rels", _ROOT_RELS)
            z.writestr("word/document.xml", document)
        docx.write_bytes(buf.getvalue())
        html = pypandoc.convert_file(str(docx), "html", format="docx", extra_args=["--mathml"])
    parts = re.split(r"@@EQ(\d+)@@", html)
    out = [""] * len(maths)
    for i in range(1, len(parts), 2):
        out[int(parts[i])] = "".join(re.findall(r"<math\b.*?</math>", parts[i + 1], re.S))
    return out
